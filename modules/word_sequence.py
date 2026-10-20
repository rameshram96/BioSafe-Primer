"""
Appendix A for the Word report: the vector sequence in text form, primers
coloured (FP blue, RP orange) and one real Word comment per primer.

Requires python-docx >= 1.2 for comments (Document.add_comment). With an older
version the appendix is still produced (coloured sequence) and
the caller is told comments were skipped.
"""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .seq_annot import (BLOCK, FP_COLOR, PER_LINE, RP_COLOR,
                        build_cover as _build_cover, color_for as _color_for,
                        verify_footprint)

MONO = 'Courier New'
LABEL_W = 7          # width of the position label, in characters
FONT_PT = 8.5


def comments_supported(doc):
    return hasattr(doc, 'add_comment')


def _style_run(run, size=FONT_PT, color=None, bold=False):
    run.font.name = MONO
    run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), MONO)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _pos_text(f, seq_len):
    return ', '.join(f"{s + 1}–{e}" for s, e in f['spans'])


def _comment_for(doc, f, j, sequence, author, initials, runs):
    ok = verify_footprint(f, sequence)
    lines = []
    amp = f['name'].rsplit('_', 1)[0]
    lines.append(f"{f['name']}  ({'Forward' if f['kind'] == 'FP' else 'Reverse'} primer, {amp})")
    lines.append(f"Sequence (5'→3', as ordered): {f['sequence']}")
    lines.append(f"Length {f['length']} nt · Tm {f['tm']} °C · GC {f['gc']} %")
    if f['kind'] == 'RP':
        lines.append("Binds the bottom strand: the highlighted top-strand text is the "
                     "reverse complement of the ordered sequence.")
    n = len(f['spans'])
    if n > 1:
        lines.append(f"Spans the plasmid origin — this is segment {j + 1} of {n} "
                     f"(position {f['spans'][j][0] + 1}–{f['spans'][j][1]}); "
                     f"full footprint {_pos_text(f, 0)}.")
    else:
        lines.append(f"Position {f['spans'][0][0] + 1}–{f['spans'][0][1]} (1-based).")
    lines.append("Sequence check vs vector: " + ("match ✓" if ok else "MISMATCH ✗ — verify"))

    c = doc.add_comment(runs=runs, text=lines[0], author=author, initials=initials)
    for ln in lines[1:]:
        c.add_paragraph(ln)
    return ok


def add_sequence_appendix(doc, rd, with_comments=True, author='SOPrimeR'):
    """
    Append 'Appendix A' to `doc`. Returns a small status dict:
    {'comments': n_added, 'comments_skipped_reason': str|None, 'mismatches': [names]}.
    """
    sequence = rd['sequence'].upper()
    seq_len = len(sequence)
    fps = rd['footprints']
    status = {'comments': 0, 'comments_skipped_reason': None, 'mismatches': []}

    can_comment = with_comments and comments_supported(doc)
    if with_comments and not can_comment:
        status['comments_skipped_reason'] = "python-docx >= 1.2 is required for Word comments"

    doc.add_page_break()
    h = doc.add_heading("Appendix A. Vector Sequence with Annotated Primers", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor.from_string('1A237E')

    p = doc.add_paragraph()
    r = p.add_run(f"{rd['meta']['vector_name']} — {seq_len:,} bp, circular. "
                  "Numbering is 1-based, 100 bp per line, in blocks of 10. ")
    r.font.size = Pt(9)
    for label, col in (("Forward primer", FP_COLOR), ("Reverse primer", RP_COLOR)):
        run = p.add_run(f" {label} ")
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(col)
        p.add_run("  ").font.size = Pt(9)
    note = ("Each primer carries a Word comment with its name, sequence, Tm, GC% and "
            "position (open the Review pane to read them). Reverse primers are shown as "
            "their top-strand footprint; the comment gives the ordered sequence."
            if can_comment else
            "Primers are shown in colour (forward blue, reverse orange).")
    n = doc.add_paragraph()
    nr = n.add_run(note)
    nr.italic = True
    nr.font.size = Pt(8.5)

    cover = _build_cover(seq_len, fps)
    first_last = {}                       # key -> [first_run, last_run]
    initials = ''.join(w[0] for w in author.split()[:2]).upper() or 'SP'

    for line_start in range(0, seq_len, PER_LINE):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.keep_together = True

        lab = para.add_run(f"{line_start + 1:>{LABEL_W}}  ")
        _style_run(lab, color='78909C')

        # (text, keyset) segments, merged while the keyset is unchanged
        segs = []

        def push(ch, keys):
            if segs and segs[-1][1] == keys:
                segs[-1][0] += ch
            else:
                segs.append([ch, keys])

        line_end = min(line_start + PER_LINE, seq_len)
        for pos in range(line_start, line_end):
            if pos > line_start and (pos - line_start) % BLOCK == 0:
                prev, cur = cover[pos - 1], cover[pos]
                push(' ', frozenset(prev) if (prev and prev == cur) else frozenset())
            push(sequence[pos], frozenset(cover[pos]))

        for text, keys in segs:
            run = para.add_run(text)
            _style_run(run)
            col = _color_for(keys, fps)
            if col:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(col)
                for k in keys:
                    if k not in first_last:
                        first_last[k] = [run, run]
                    else:
                        first_last[k][1] = run

    if can_comment:
        for (i, j), (r1, r2) in sorted(first_last.items()):
            runs = r1 if r1 is r2 else [r1, r2]
            ok = _comment_for(doc, fps[i], j, sequence, author, initials, runs)
            status['comments'] += 1
            if not ok and fps[i]['name'] not in status['mismatches']:
                status['mismatches'].append(fps[i]['name'])
    else:
        for f in fps:
            if not verify_footprint(f, sequence):
                status['mismatches'].append(f['name'])

    if status['mismatches']:
        warn = doc.add_paragraph()
        wr = warn.add_run("Warning: these primers do not match the vector sequence at their recorded "
                          "position - " + ", ".join(status['mismatches']) + ". Verify before submission.")
        wr.italic = True
        wr.font.size = Pt(9)
        wr.font.color.rgb = RGBColor.from_string('B71C1C')
    return status
